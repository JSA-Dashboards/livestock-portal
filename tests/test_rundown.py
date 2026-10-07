"""
Pin the Cattle Market Rundown slide.

The slide goes to clients with no human between the feed and the deck, so the
tests here are about the ways it can be WRONG WITHOUT LOOKING WRONG:

  - a figure quietly disagreeing with the letter built from the same ctx
  - a bullet silently dropped, which on a slide is a statement
  - the 5-day average quoting the wrong window
  - text printed through the logo, which renders without complaint
  - bullets that report the right indent level through python-pptx and have
    no bullet glyph at all

The fixture is the REAL 2026-10-05 data -- the day of the hand-typed slide
this replaced -- so the expected strings are checkable against a deck Ross
actually sent rather than against numbers invented to agree.

    python -m pytest tests/test_rundown.py -q
"""
from __future__ import annotations

import io
from datetime import date
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parent.parent
import sys  # noqa: E402

if str(REPO) not in sys.path:
    sys.path.insert(0, str(REPO))

from letter import rundown  # noqa: E402


# The figures USDA served on Monday 2026-10-05, verified live that week.
CTX = {
    "cash": {
        "live": {"this_week": 220.03, "last_week": 220.73},
        "dressed": {"this_week": 346.35, "last_week": 347.93},
        "volume": {"confirmed": 88019.0, "confirmed_last_week": 47138.0,
                   "d14": 70390.0, "d14_last_week": 36643.0,
                   "d30": 17629.0, "d30_last_week": 10495.0,
                   "report_date": "2026-10-05"},
        "report_date": "2026-10-05",
    },
    "cutout": {
        "choice": {"value": 378.26, "change": 4.07, "avg5": 379.38},
        "select": {"value": 358.57, "change": 4.08, "avg5": 358.14},
        "report_date": "2026-10-05",
        "grading": {"pct": 87.3, "pct_last_week": 88.0},
    },
    "carcass_weights": {
        "value": 891.0, "week_ending": "2026-09-19",
        "last_week": 893.0, "year_ago": 881.0,
        "steers": {"value": 979.0, "week_ending": "2026-09-19",
                   "last_week": 977.0, "year_ago": 967.0},
    },
    # CME had published 10/02 at 337.76 when this slide was built; 337.66
    # was our forward estimate for 10/05, the row still filling.
    "fci": {"value": 337.66, "date": "2026-10-05",
            "published": {"value": 337.76, "date": "2026-10-02",
                          "change": 0.92}},
    "slaughter": {"weekly": {"ytd_chg_pct": -7.5},
                  "beef_production": {"ytd_chg_pct": -5.2}},
}

EXPECTED = [
    (0, "Last week's cash trade"),
    (1, "220.03 live versus 220.73 week earlier"),
    (1, "346.35 dressed versus 347.93 week earlier"),
    (1, "Volume: +40,881/hd WoW"),
    (2, "88,019 head total"),
    (3, "1-14-day: 70,390. +33,747/hd WoW"),
    (3, "15-30 day: 17,629. +7,134/hd WoW"),
    (0, "Actual carcass weights as of 9/19/26- 891# versus 893# LW and 881# LY"),
    (1, "Steers- 979# versus 977# LW and 967# LY"),
    (0, "Choice- 378.26. +4.07 Monday PM."),
    (1, "5-day average- 379.38."),
    (0, "Select- 358.57. +4.08 Monday PM."),
    (1, "5-day average- 358.14."),
    (0, "Choice/Select spread: 19.69"),
    (1, "Choice & Higher- 87.3%"),
    (0, "Feeder Index: 337.76"),
    (0, "Slaughter -7.5% YTD"),
    (0, "Beef Production -5.2% YTD"),
]


def test_it_reproduces_the_slide_that_was_typed_by_hand():
    """
    Line for line against the 2026-10-05 deck, including the punctuation.

    The volume lines are the ones worth having pinned: every figure on them
    is a SUBTRACTION the slide shows only the result of, so a sign slip or a
    swapped week reads as a perfectly plausible market.
    """
    assert rundown.rows(CTX) == EXPECTED


def test_the_index_is_cme_s_published_figure_not_the_forward_estimate():
    """
    `fci["value"]` is the index CME will print NEXT, so on an afternoon
    slide it is a part-day estimate of a day that has not finished. On
    2026-10-07 it read 335.86 for index date 10/07 -- 228 locations against
    the previous day's 277 -- where CME had published 10/06 at 337.87.
    Ross caught it on the letter; the slide had the same fault.
    """
    line = [t for _, t in rundown.rows(CTX) if t.startswith("Feeder Index")][0]
    assert line == "Feeder Index: 337.76"
    assert "337.66" not in line, "still quoting the forward estimate"


def test_the_line_is_bare_with_no_date_and_no_label():
    """
    A first attempt printed "337.76 (10/2)" and, on the fallback,
    "(JSA est. 10/7)". Ross rejected both: the line is the actual CME
    feeder index and a parenthesis on it is noise or an excuse.
    """
    line = [t for _, t in rundown.rows(CTX) if t.startswith("Feeder Index")][0]
    assert "(" not in line and "est" not in line.lower()


def test_it_never_falls_back_to_the_forward_estimate():
    """
    THE FALLBACK WAS THE BUG. The 2026-10-07 fix fell back to fci["value"]
    when the published block was missing, which on the deployed app -- a
    different Snowflake, where the CME table did not answer -- printed the
    exact figure Ross had reported, with a label on it. Marking beats
    reproducing the defect.
    """
    ctx = {**CTX, "fci": {"value": 335.86, "date": "2026-10-07"}}
    line = [t for _, t in rundown.rows(ctx) if t.startswith("Feeder Index")][0]
    assert "335.86" not in line, "fell back to the forward estimate"
    assert rundown.MISSING in line


def test_the_slide_and_the_letter_quote_the_same_index():
    """
    CLAUDE.md records twice that the letter and a dashboard quoted the same
    figure and disagreed, each defensible and neither raising. Both now read
    the same `published` block, so pin that they agree.
    """
    from letter import render
    line = [t for _, t in rundown.rows(CTX) if t.startswith("Feeder Index")][0]
    assert render.fci_published(CTX["fci"]) in line


def test_a_missing_figure_is_marked_and_the_bullet_still_prints():
    """
    A BULLET THAT VANISHES IS A STATEMENT, and a false one -- it reads as
    "nothing to report there". Every row survives an empty ctx; the holes are
    visible instead.
    """
    rows = rundown.rows({})
    assert len(rows) == len(EXPECTED)
    assert [lvl for lvl, _ in rows] == [lvl for lvl, _ in EXPECTED]
    assert rundown.missing(rows) > 0
    assert all(rundown.MISSING in text or text == "Last week's cash trade"
               for _, text in rows)


def test_the_marker_is_plain_text_not_the_letter_s_html():
    """render.MISSING is wrapped in a <span> for the letter's CSS. Inside a
    PowerPoint text box that markup would print literally."""
    assert "<" not in rundown.MISSING and ">" not in rundown.MISSING


def test_the_week_ending_is_printed_from_the_data():
    """
    The hand-typed slide read "as of 8/19/26" for a week ending 9/19/26. AMS
    3658 runs about a fortnight behind, so a four-week-old weight looks no
    different from a two-week-old one and the typo survived. Printing the
    date the data carries is the fix; this pins that it comes from the data.
    """
    ctx = {**CTX, "carcass_weights": {**CTX["carcass_weights"],
                                      "week_ending": "2026-08-19"}}
    line = [t for _, t in rundown.rows(ctx) if t.startswith("Actual carcass")][0]
    assert "8/19/26" in line


def test_the_spread_is_derived_so_the_three_numbers_agree():
    """Choice, Select and the spread are on the same slide. Taking the spread
    from anywhere but the two values printed above it invites a slide whose
    own arithmetic does not check out."""
    ctx = {**CTX, "cutout": {**CTX["cutout"],
                             "choice": {"value": 400.0, "change": 1.0, "avg5": 1.0},
                             "select": {"value": 375.5, "change": 1.0, "avg5": 1.0}}}
    assert (0, "Choice/Select spread: 24.50") in rundown.rows(ctx)


def test_the_session_label_follows_the_report_date():
    """'Monday PM' is read off the cutout's own date, not from today."""
    ctx = {**CTX, "cutout": {**CTX["cutout"], "report_date": "2026-10-08"}}
    assert any("Thursday PM." in t for _, t in rundown.rows(ctx))


# -- freshness ---------------------------------------------------------------

def test_a_previous_session_s_cutout_is_called_out():
    """
    The failure this exists for is silent: Friday's cutout is a perfectly good
    number and looks exactly like Monday's. Nothing rebuilds the slide at 3pm,
    so the page has to say which session it is holding.
    """
    f = rundown.freshness(CTX, today=date(2026, 10, 6))
    assert f["stale"] is True
    assert "10/5/26" in f["message"]


def test_today_s_cutout_is_not_called_stale():
    f = rundown.freshness(CTX, today=date(2026, 10, 5))
    assert f["stale"] is False


def test_a_weekend_expects_friday_not_saturday():
    """
    Saturday has no session of its own, so holding FRIDAY's cutout is current
    and must not be flagged -- a banner that cries stale every weekend trains
    the reader to ignore it on the Monday it matters.

    Holding something older than Friday on a Saturday is still stale, which
    is why this is two assertions and not one.
    """
    fri = {"cutout": {"report_date": "2026-10-09"}}
    assert rundown.freshness(fri, today=date(2026, 10, 10))["stale"] is False
    assert rundown.freshness(CTX, today=date(2026, 10, 10))["stale"] is True


def test_no_cutout_at_all_is_stale_and_says_so():
    f = rundown.freshness({}, today=date(2026, 10, 6))
    assert f["stale"] is True
    assert "no cutout" in f["message"].lower()


# -- the deck ----------------------------------------------------------------

def test_the_text_cannot_print_through_the_logo():
    """
    IT DID. At 20/17/15/14pt the eighteen rows ran to 6.60in and the logo sat
    at 6.55in, so "Beef Production -5.2% YTD" printed through the wordmark.
    python-pptx reported a perfectly well-formed file; only a picture of the
    slide showed it. Arithmetic here means nobody has to take that picture --
    which matters, because taking it is what cost Ross an unsaved deck.
    """
    bottom = rundown.BODY_TOP_IN + rundown.text_height_in(rundown.rows(CTX))
    assert bottom < rundown.LOGO_TOP_IN, (
        f"text runs to {bottom:.2f}in, logo starts at {rundown.LOGO_TOP_IN}in")
    assert rundown.LOGO_TOP_IN + rundown.LOGO_H_IN <= 7.5


def test_every_paragraph_gets_a_real_bullet_glyph():
    """
    SETTING para.level IS NOT ENOUGH. A level selects a style from the
    layout's list styles and a blank-layout text box has none, so the first
    build was flat, unindented, bullet-less text that still reported the
    right `level` back through python-pptx. Assert the XML, not the API.
    """
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.oxml.ns import qn

    buf = io.BytesIO()
    rundown.build_pptx(rundown.rows(CTX), buf)
    buf.seek(0)
    body = [s for s in Presentation(buf).slides[0].shapes if s.has_text_frame][1]
    paras = body.text_frame.paragraphs
    assert len(paras) == len(EXPECTED)
    for p in paras:
        pPr = p._p.find(qn("a:pPr"))
        assert pPr is not None and pPr.find(qn("a:buChar")) is not None, \
            f"no bullet on {p.text!r}"
        assert pPr.get("marL") and pPr.get("indent")


def _body_paragraphs(logo=None, agmarket=None):
    from pptx import Presentation
    buf = io.BytesIO()
    rundown.build_pptx(rundown.rows(CTX), buf, logo=logo, agmarket=agmarket)
    buf.seek(0)
    prs = Presentation(buf)
    shapes = prs.slides[0].shapes
    body = [s for s in shapes if s.has_text_frame][1]
    return prs, shapes, body.text_frame.paragraphs


def test_every_run_is_aptos_sixteen():
    """
    ONE FONT, ONE SIZE -- Ross's house format, specified 2026-10-06. An
    earlier version tapered 19/16/14/13 by indent level, which is what the
    deck does by default and is not what this slide uses.
    """
    pytest.importorskip("pptx")
    from pptx.util import Pt
    _, _, paras = _body_paragraphs()
    for p in paras:
        for r in p.runs:
            assert r.font.name == "Aptos", f"{r.text!r} is {r.font.name}"
            assert r.font.size == Pt(16), f"{r.text!r} is {r.font.size}"


def test_the_spacing_is_identical_on_every_paragraph():
    """
    "spacing- everything needs to be the same". PowerPoint's default puts
    extra space BEFORE a top-level bullet, so the gap above "Choice-" came
    out larger than the gap above "5-day average-". space_before is pinned
    to zero as well as space_after, because leaving it unset is what let the
    deck default back in.
    """
    pytest.importorskip("pptx")
    from pptx.util import Pt
    _, _, paras = _body_paragraphs()
    assert {p.space_after for p in paras} == {Pt(rundown.SPACE_AFTER_PT)}
    assert {p.space_before for p in paras} == {Pt(0)}


def test_only_the_heading_is_underlined():
    pytest.importorskip("pptx")
    _, _, paras = _body_paragraphs()
    underlined = [p.text for p in paras
                  if any(r.font.underline for r in p.runs)]
    assert underlined == [rundown.HEADING]


def test_both_marks_sit_on_the_slide_without_overlapping():
    """
    John Stewart bottom-left, AgMarket.Net bottom-right. They are pulled from
    the real deck rather than cropped from a screenshot, so they are the
    full-resolution originals.
    """
    pytest.importorskip("pptx")
    jsa = REPO / "assets" / "logo-full.png"
    agm = REPO / "assets" / "agmarket-net.png"
    assert jsa.exists() and agm.exists()
    prs, shapes, _ = _body_paragraphs(logo=jsa, agmarket=agm)
    pics = sorted((s for s in shapes if s.shape_type == 13),
                  key=lambda s: s.left)
    assert len(pics) == 2, "expected both marks"
    left, right = pics
    assert left.left < right.left
    assert left.left + left.width < right.left, "the two marks overlap"
    assert right.left + right.width <= prs.slide_width, "AgMarket runs off the sheet"
    for pic in pics:
        assert pic.top + pic.height <= prs.slide_height


def test_the_deck_is_sixteen_by_nine():
    """python-pptx's default template is 4:3 and would letterbox inside the
    weekly deck."""
    pytest.importorskip("pptx")
    from pptx import Presentation
    buf = io.BytesIO()
    rundown.build_pptx(rundown.rows(CTX), buf)
    buf.seek(0)
    prs = Presentation(buf)
    assert round(prs.slide_width.inches, 2) == 13.33
    assert round(prs.slide_height.inches, 2) == 7.50


def test_the_slide_carries_every_figure_as_text():
    """A deck that renders but has lost a number is the expensive failure."""
    pytest.importorskip("pptx")
    from pptx import Presentation
    buf = io.BytesIO()
    rundown.build_pptx(rundown.rows(CTX), buf)
    buf.seek(0)
    text = "\n".join(s.text_frame.text for s in Presentation(buf).slides[0].shapes
                     if s.has_text_frame)
    for probe in ("88,019", "+33,747", "379.38", "19.69", "87.3%", "337.76",
                  "-7.5% YTD", "-5.2% YTD", "979#", "9/19/26"):
        assert probe in text, probe


# -- the rule that keeps it honest -------------------------------------------

def test_the_module_cannot_fetch_anything():
    """
    IT MUST BUILD FROM ctx AND NOTHING ELSE.

    CLAUDE.md records twice that the letter and a dashboard quoted the same
    figure and disagreed, each defensible in isolation and neither raising. A
    slide that did its own fetching would be a third number in that argument.
    Reading the source rather than the imports catches a late `import
    requests` inside a function, which is how that rule usually gets broken.
    """
    import ast

    # AST, NOT A SUBSTRING SEARCH. The first version of this test grepped for
    # "requests" and failed on the module's own docstring, which explains why
    # it must not import it. Prose about a rule is not a breach of it; an
    # import node is.
    tree = ast.parse((REPO / "letter" / "rundown.py").read_text(encoding="utf-8"))
    banned = {"requests", "urllib", "urllib3", "httpx", "psycopg2",
              "snowflake", "snowflake_db", "sources"}
    found = set()
    for node in ast.walk(tree):          # ast.walk reaches imports inside functions
        if isinstance(node, ast.Import):
            found |= {a.name.split(".")[0] for a in node.names}
        elif isinstance(node, ast.ImportFrom):
            found.add((node.module or "").split(".")[0])
            if node.level and not node.module:
                found |= {a.name for a in node.names}
            elif node.level:
                found |= {a.name for a in node.names}
    assert not (found & banned), f"rundown.py imports {sorted(found & banned)}"


def test_the_five_day_average_window_is_documented_where_it_is_computed():
    """
    The convention lives in sources.fetch_cutout and the slide only prints it.
    It changed on 2026-10-06 from including the day being reported to
    excluding it -- 378.94 against 379.38 for 2026-10-05 -- and the letter had
    been quoting the other one. Pin the window so it cannot drift back
    silently.
    """
    src = (REPO / "letter" / "sources.py").read_text(encoding="utf-8")
    assert "iloc[-6:-1]" in src, "the 5-day average no longer excludes today"
    assert "tail(5).mean()" not in src
