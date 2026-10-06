"""
Pin the Sterling Beef Profit Tracker parser and the slide it feeds.

THE FIXTURE IS SYNTHETIC, AND DELIBERATELY SO. This repository is public and
Sterling is a paid subscription whose PDF says "considered proprietary
material". The sample below copies the LAYOUT of a real tracker -- the row
order, the label wording, the footnote markers, the accounting parentheses,
the extraction quirks -- and none of its figures. Every number here is made
up. A fixture of Sterling's real output would publish their product.

What the real tracker is checked against is the live fetch, not a committed
copy; the tests below prove the parser handles the SHAPE, which is what
breaks when Sterling changes their layout.

    python -m pytest tests/test_sterling.py -q
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parent.parent
if str(REPO) not in sys.path:
    sys.path.insert(0, str(REPO))

from letter import rundown, sterling  # noqa: E402


# Same shape as page 1 of a real tracker. Invented figures.
SAMPLE = """Week Ending  Week Ago Month Ago Year Ago
January 3, 2027
  Feedlot Margin - Unhedged ($ / head) ($111.11) ($222.22) ($333.33) $444.44
  Choice Steers (5-Area Direct, $ / cwt) 200.00 201.00 202.00 203.00
  Feeder Steer (Ok City 750-800 lb, $ / cwt)
For this Week's Feedlot Placement 300.00 301.00 302.00 303.00
  Packer Margin ($ / head) $11.11 $22.22 $33.33 ($44.44)
  Beef Cutout 1  ($ / cwt) 370.00 371.00 372.00 373.00
  Drop Credit 2 ($/hd) 230.00 231.00 232.00 233.00
January 3, 2027 Week Ago Month Ago Year Ago
  Cattle Slaughter 500,000 510,000 520,000 530,000
             Steer & Heifer 400,000 410,000 420,000 430,000
                     Fed Plant Capacity Utilization 70.1% 71.2% 72.3% 73.4%
             Cows 90,000 0 92,000 93,000
                     Cow Plant Capacity Utilization 50.1% 51.2% 52.3% 53.4%
  Beef Production 400.1 401.2 402.3 403.4
(federally inspected, mil. lbs.)
              Carcass Weight - average - all cattle (lbs.) 800 801 802 803
  Relative Feeding Cost (against current Placement week)
  Feeder Steer 70.00% 71.00% 72.00% 73.00%
Annual Projections - Sept. 14, 2026 2027* 2026 2025 2024
   Cow-Calf Margin 3($ / cow) 1000.00 900.00 500.00 300.00
    Estimated annual revenue - annual variable costs
   Feedlot Margin ($ / head) 20.00 490.00 110.00 210.00
   Packer Margin ($ / head) ($150.00) ($130.00) ($70) 90.00
Sterling Beef Profit Tracker
© Sterling Marketing 1991-2027
"""


@pytest.fixture(scope="module")
def parsed():
    return sterling.parse(SAMPLE)


def test_the_week_ending_is_not_the_first_line(parsed):
    """
    pypdf puts the column header above the date, so lines[0] is
    "Week Ending  Week Ago Month Ago Year Ago" and taking it as the date
    yields None -- which is exactly what the first version did.
    """
    assert parsed["week_ending"].isoformat() == "2027-01-03"


def test_sterling_s_own_month_abbreviation_parses():
    """"Sept." is not a month name strptime accepts under any locale."""
    assert sterling._date("Annual Projections - Sept. 14, 2026").isoformat() == "2026-09-14"
    assert sterling._date("September 26, 2026").isoformat() == "2026-09-26"


def test_parentheses_are_negative(parsed):
    """
    Sterling writes losses in accounting style. Reading ($111.11) as positive
    turns the worst feedlot margin in two years into a profit -- a wrong
    number that looks entirely plausible on a slide.
    """
    assert parsed["weekly"]["feedlot_margin"] == [-111.11, -222.22, -333.33, 444.44]
    assert parsed["weekly"]["packer_margin"] == [11.11, 22.22, 33.33, -44.44]
    assert parsed["annual"]["packer_margin"] == [-150.0, -130.0, -70.0, 90.0]


def test_footnote_digits_in_a_label_are_not_read_as_data(parsed):
    """
    "Beef Cutout 1 ($ / cwt) 370.00 ..." and "Cow-Calf Margin 3($ / cow) ..."
    both carry a footnote marker before the figures. Taking the FIRST four
    numbers would read the marker as a price and shift every column.
    """
    assert parsed["annual"]["cow_calf_margin"] == [1000.0, 900.0, 500.0, 300.0]
    assert sterling._last_four("  Beef Cutout 1  ($ / cwt) 370.00 371.00 372.00 373.00") \
        == [370.0, 371.0, 372.0, 373.0]


def test_a_short_row_is_refused_rather_than_padded():
    """
    A padded row would still render and would shift every column silently.
    None makes a layout change show up as a gap.
    """
    assert sterling._last_four("  Cattle Slaughter 500,000 510,000") is None
    assert sterling._last_four("  Feeder Steer (Ok City 750-800 lb, $ / cwt)") is None


def test_the_annual_block_does_not_steal_the_weekly_labels(parsed):
    """
    "Feedlot Margin" and "Packer Margin" appear in BOTH tables. Parsed by
    label alone the annual figures would overwrite the weekly ones, or the
    reverse, depending on line order.
    """
    assert parsed["weekly"]["feedlot_margin"][0] == -111.11
    assert parsed["annual"]["feedlot_margin"][0] == 20.0


def test_the_annual_header_year_is_not_counted_as_a_column(parsed):
    """
    "Annual Projections - Sept. 14, 2026 2027* 2026 2025 2024" has FIVE
    four-digit years on it. Scanning the whole line yields a duplicate and
    silently shifts the header by one column.
    """
    assert parsed["annual_years"] == ["2027", "2026", "2025", "2024"]


def test_sterling_s_own_label_is_kept_verbatim(parsed):
    """They write "Sept."; strftime gives "Sep". Reformatting their header is
    the same class of mistake as recomputing their figures."""
    assert parsed["annual_label"] == "Annual Projections - Sept. 14, 2026"


def test_a_zero_is_reproduced_and_flagged_not_repaired(parsed):
    """
    The real 2026-09-28 tracker reports Cows week-ago as 0 while the capacity
    row beneath reads 60.0% for the same column. It is wrong in Sterling's
    PDF. Repairing it would mean inventing a figure and attributing it to
    them, so it is reproduced and reported instead.
    """
    assert parsed["weekly"]["cows"][1] == 0
    notes = sterling.anomalies(parsed)
    assert any("Cows as 0" in n for n in notes)
    assert any("capacity utilisation" in n.lower() for n in notes)


def test_the_parser_never_reaches_for_ams_to_second_guess_sterling():
    """
    AMS has its own weekly split and it disagrees -- 382,000/102,000 against
    Sterling's 395,428/82,280 for the same week. Blending them would produce
    a table that is neither source's. The module may import mailbox and pypdf
    to GET the PDF; it must not import a USDA fetcher.
    """
    import ast
    tree = ast.parse((REPO / "letter" / "sterling.py").read_text(encoding="utf-8"))
    mods = set()
    for n in ast.walk(tree):
        if isinstance(n, ast.Import):
            mods |= {a.name.split(".")[0] for a in n.names}
        elif isinstance(n, ast.ImportFrom):
            mods.add((n.module or "").split(".")[0])
            if n.level:
                mods |= {a.name for a in n.names}
    assert "sources" not in mods, "sterling.py reaches for the USDA fetchers"


# -- the slide ---------------------------------------------------------------

def test_the_margins_block_reads_like_the_hand_built_slide(parsed):
    rows = rundown.sterling_rows(parsed)
    assert rows[0] == (0, "Feedlot Margins-")
    assert rows[1] == (1, "Current- (111.11) (Unhedged)")
    assert rows[5] == (0, "Packer Margins-")
    assert rows[6] == (1, "Current- 11.11")
    assert rows[9] == (1, "Last year- (44.44)")


def test_both_margin_blocks_say_last_year_the_same_way(parsed):
    """The hand-built slide writes "Last year-" under Feedlot and "Last Year-"
    under Packer. That is a slip, not a distinction."""
    rows = [t for _, t in rundown.sterling_rows(parsed)]
    assert sum(t.startswith("Last year-") for t in rows) == 2
    assert not any(t.startswith("Last Year-") for t in rows)


def test_the_slide_carries_the_attribution(parsed):
    """Sterling's data, Sterling's credit. Not optional."""
    pytest.importorskip("pptx")
    from pptx import Presentation
    buf = io.BytesIO()
    rundown.build_sterling_pptx(parsed, buf)
    buf.seek(0)
    text = "\n".join(s.text_frame.text for s in Presentation(buf).slides[0].shapes
                     if s.has_text_frame)
    assert sterling.ATTRIBUTION in text


def test_the_tables_are_real_tables_not_pictures(parsed):
    """
    The hand-built slide pastes them as screenshots -- image9.png is 603x102
    stretched to 6.28in, which is why it is soft on a projector. Native
    tables stay sharp and keep the numbers selectable.
    """
    pytest.importorskip("pptx")
    from pptx import Presentation
    buf = io.BytesIO()
    rundown.build_sterling_pptx(parsed, buf)
    buf.seek(0)
    shapes = Presentation(buf).slides[0].shapes
    tables = [s for s in shapes if getattr(s, "has_table", False) and s.has_table]
    assert len(tables) == 2
    assert (len(tables[0].table.rows), len(tables[0].table.columns)) == (6, 5)
    assert (len(tables[1].table.rows), len(tables[1].table.columns)) == (4, 5)


def test_every_sterling_figure_reaches_the_slide(parsed):
    pytest.importorskip("pptx")
    from pptx import Presentation
    buf = io.BytesIO()
    rundown.build_sterling_pptx(parsed, buf)
    buf.seek(0)
    shapes = Presentation(buf).slides[0].shapes
    text = "\n".join(s.text_frame.text for s in shapes if s.has_text_frame)
    for t in (s.table for s in shapes if getattr(s, "has_table", False) and s.has_table):
        text += "\n" + "\n".join(c.text for r in t.rows for c in r.cells)
    for probe in ("500,000", "400,000", "70.1%", "90,000", "50.1%",
                  "January 3, 2027", "2027*", "1,000.00", "(150.00)",
                  "(111.11)", "11.11"):
        assert probe in text, probe


def test_the_slide_sits_where_the_deck_puts_it(parsed):
    """Geometry measured off slide 7 of the real deck, so a generated slide
    drops in without nudging."""
    pytest.importorskip("pptx")
    from pptx import Presentation
    buf = io.BytesIO()
    rundown.build_sterling_pptx(parsed, buf)
    buf.seek(0)
    prs = Presentation(buf)
    assert round(prs.slide_width.inches, 2) == 13.33
    tables = sorted((s for s in prs.slides[0].shapes
                     if getattr(s, "has_table", False) and s.has_table),
                    key=lambda s: s.top)
    assert round(tables[0].top.inches, 2) == 3.92
    assert round(tables[1].top.inches, 2) == 5.19
